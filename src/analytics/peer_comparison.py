"""
Sprint 3 - Day 20 Peer Comparison Excel

Creates:
    output/peer_comparison.xlsx

Features:
- 11 peer-group sheets
- Company financial metric values
- Percentile ranks
- Benchmark identification
- Peer median
- Percentile-based cell formatting
"""

from __future__ import annotations

import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[2]

DB_PATH = PROJECT_ROOT / "db" / "nifty100.db"
PEER_GROUPS_PATH = PROJECT_ROOT / "data" / "source" / "peer_groups.xlsx"
OUTPUT_PATH = PROJECT_ROOT / "output" / "peer_comparison.xlsx"


METRICS = [
    ("ROE", "return_on_equity_pct"),
    ("ROCE", "return_on_capital_employed_pct"),
    ("NPM", "net_profit_margin_pct"),
    ("D/E", "debt_to_equity"),
    ("FCF", "free_cash_flow_cr"),
    ("PAT CAGR 5Y", "pat_cagr_5yr"),
    ("Revenue CAGR 5Y", "revenue_cagr_5yr"),
    ("EPS CAGR 5Y", "eps_cagr_5yr"),
    ("Interest Coverage", "interest_coverage"),
    ("Asset Turnover", "asset_turnover"),
]


def load_peer_groups() -> pd.DataFrame:
    """Load peer-group membership from Excel."""

    df = pd.read_excel(PEER_GROUPS_PATH)

    required = {
        "peer_group_name",
        "company_id",
        "is_benchmark",
    }

    missing = required - set(df.columns)

    if missing:
        raise ValueError(
            f"Missing peer-group columns: {sorted(missing)}"
        )

    df = df[
        [
            "peer_group_name",
            "company_id",
            "is_benchmark",
        ]
    ].copy()

    df["company_id"] = (
        df["company_id"]
        .astype(str)
        .str.strip()
    )

    df["peer_group_name"] = (
        df["peer_group_name"]
        .astype(str)
        .str.strip()
    )

    return df


def load_latest_ratios() -> pd.DataFrame:
    """Load latest financial ratios."""

    conn = sqlite3.connect(DB_PATH)

    columns = ", ".join(
        column
        for _, column in METRICS
    )

    query = f"""
        WITH latest AS (
            SELECT
                company_id,
                MAX(year) AS year
            FROM financial_ratios
            GROUP BY company_id
        )
        SELECT
            fr.company_id,
            fr.year,
            {columns}
        FROM financial_ratios fr
        INNER JOIN latest l
            ON fr.company_id = l.company_id
            AND fr.year = l.year
    """

    df = pd.read_sql_query(
        query,
        conn,
    )

    conn.close()

    return df


def load_percentiles() -> pd.DataFrame:
    """Load Day 18 percentile results."""

    conn = sqlite3.connect(DB_PATH)

    df = pd.read_sql_query(
        """
        SELECT
            company_id,
            peer_group_name,
            metric,
            percentile_rank
        FROM peer_percentiles
        """,
        conn,
    )

    conn.close()

    return df


def build_peer_sheet(
    peer_group_name: str,
    members: pd.DataFrame,
    ratios: pd.DataFrame,
    percentiles: pd.DataFrame,
) -> pd.DataFrame:
    """Build one peer-group comparison DataFrame."""

    group = members[
        members["peer_group_name"] == peer_group_name
    ].copy()

    result = group.merge(
        ratios,
        on="company_id",
        how="left",
    )

    result = result.rename(
        columns={
            "company_id": "Company ID",
            "is_benchmark": "Benchmark",
            "year": "Year",
        }
    )

    # Convert percentile data from long to wide.
    pct = percentiles[
        percentiles["peer_group_name"]
        == peer_group_name
    ].copy()

    if not pct.empty:
        pct_wide = pct.pivot_table(
            index="company_id",
            columns="metric",
            values="percentile_rank",
            aggfunc="first",
        ).reset_index()

        pct_wide.columns = [
            (
                "Company ID"
                if column == "company_id"
                else f"{column} Percentile"
            )
            for column in pct_wide.columns
        ]

        result = result.merge(
            pct_wide,
            on="Company ID",
            how="left",
        )

    # Add peer median columns.
    for metric_name, source_column in METRICS:

        if source_column in result.columns:
            median_value = pd.to_numeric(
                result[source_column],
                errors="coerce",
            ).median()

            result[
                f"{metric_name} Peer Median"
            ] = median_value

    # Keep the required comparison structure readable.
    ordered_columns = [
        "Company ID",
        "Benchmark",
        "Year",
    ]

    for metric_name, source_column in METRICS:
        if source_column in result.columns:
            ordered_columns.append(source_column)

        percentile_column = (
            f"{metric_name} Percentile"
        )

        if percentile_column in result.columns:
            ordered_columns.append(
                percentile_column
            )

        median_column = (
            f"{metric_name} Peer Median"
        )

        if median_column in result.columns:
            ordered_columns.append(
                median_column
            )

    result = result[
        [
            column
            for column in ordered_columns
            if column in result.columns
        ]
    ]

    # Rename raw metric columns to user-friendly names.
    rename_map = {
        source: metric
        for metric, source in METRICS
    }

    result = result.rename(
        columns=rename_map
    )

    return result


def safe_sheet_name(name: str) -> str:
    """Make a valid Excel worksheet name."""

    invalid = [
        "\\",
        "/",
        "*",
        "?",
        ":",
        "[",
        "]",
    ]

    result = str(name)

    for character in invalid:
        result = result.replace(
            character,
            "_",
        )

    return result[:31]


def write_excel(
    peer_groups: pd.DataFrame,
    ratios: pd.DataFrame,
    percentiles: pd.DataFrame,
) -> int:
    """Create the peer comparison workbook."""

    OUTPUT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    peer_group_names = sorted(
        peer_groups["peer_group_name"]
        .unique()
    )

    with pd.ExcelWriter(
        OUTPUT_PATH,
        engine="openpyxl",
    ) as writer:

        for peer_group_name in peer_group_names:

            sheet_df = build_peer_sheet(
                peer_group_name,
                peer_groups,
                ratios,
                percentiles,
            )

            sheet_name = safe_sheet_name(
                peer_group_name
            )

            sheet_df.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )

    return len(peer_group_names)


def format_workbook() -> None:
    """Apply formatting and percentile color scales."""

    workbook = load_workbook(
        OUTPUT_PATH
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    benchmark_fill = PatternFill(
        fill_type="solid",
        fgColor="FFD966",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    benchmark_font = Font(
        bold=True,
    )

    for worksheet in workbook.worksheets:

        # Header formatting.
        for cell in worksheet[1]:

            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        worksheet.freeze_panes = "D2"

        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        # Benchmark highlighting.
        benchmark_column = None

        for cell in worksheet[1]:

            if cell.value == "Benchmark":
                benchmark_column = cell.column
                break

        if benchmark_column:

            for row in range(
                2,
                worksheet.max_row + 1,
            ):

                value = worksheet.cell(
                    row=row,
                    column=benchmark_column,
                ).value

                if value is True:
                    for column in range(
                        1,
                        worksheet.max_column + 1,
                    ):
                        worksheet.cell(
                            row=row,
                            column=column,
                        ).fill = benchmark_fill

                        worksheet.cell(
                            row=row,
                            column=column,
                        ).font = benchmark_font

        # Percentile color scales.
        for column in range(
            1,
            worksheet.max_column + 1,
        ):

            header = worksheet.cell(
                row=1,
                column=column,
            ).value

            if (
                isinstance(header, str)
                and header.endswith("Percentile")
            ):

                letter = get_column_letter(
                    column
                )

                if worksheet.max_row >= 2:

                    worksheet.conditional_formatting.add(
                        f"{letter}2:{letter}{worksheet.max_row}",
                        ColorScaleRule(
                            start_type="num",
                            start_value=0,
                            start_color="F8696B",
                            mid_type="num",
                            mid_value=0.5,
                            mid_color="FFEB84",
                            end_type="num",
                            end_value=1,
                            end_color="63BE7B",
                        ),
                    )

        # Number formatting.
        for row in worksheet.iter_rows(
            min_row=2,
        ):

            for cell in row:

                if isinstance(
                    cell.value,
                    float,
                ):

                    cell.number_format = "0.00"

        # Column widths.
        for column_cells in worksheet.columns:

            column_letter = (
                get_column_letter(
                    column_cells[0].column
                )
            )

            max_length = 0

            for cell in column_cells:

                value = (
                    ""
                    if cell.value is None
                    else str(cell.value)
                )

                max_length = max(
                    max_length,
                    len(value),
                )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 12),
                28,
            )

    workbook.save(
        OUTPUT_PATH
    )


def validate_output(
    peer_groups: pd.DataFrame,
) -> None:
    """Validate generated workbook."""

    workbook = load_workbook(
        OUTPUT_PATH,
        read_only=True,
    )

    print("\nVALIDATION")
    print("-" * 60)

    print(
        "Expected sheets:",
        peer_groups[
            "peer_group_name"
        ].nunique(),
    )

    print(
        "Actual sheets:",
        len(workbook.sheetnames),
    )

    print(
        "Sheet names:"
    )

    for sheet_name in workbook.sheetnames:
        print(
            " ",
            sheet_name,
        )

    workbook.close()


def main() -> None:

    print(
        "\nSPRINT 3 - DAY 20 "
        "PEER COMPARISON EXCEL\n"
    )

    peer_groups = load_peer_groups()

    ratios = load_latest_ratios()

    percentiles = load_percentiles()

    print(
        "Peer-group companies:",
        peer_groups["company_id"].nunique(),
    )

    print(
        "Peer groups:",
        peer_groups[
            "peer_group_name"
        ].nunique(),
    )

    print(
        "Latest ratio companies:",
        ratios["company_id"].nunique(),
    )

    print(
        "Percentile rows:",
        len(percentiles),
    )

    sheets_created = write_excel(
        peer_groups,
        ratios,
        percentiles,
    )

    format_workbook()

    print(
        "\nSheets created:",
        sheets_created,
    )

    print(
        "Output:",
        OUTPUT_PATH,
    )

    validate_output(
        peer_groups
    )

    print(
        "\nDAY 20 PEER COMPARISON "
        "EXCEL COMPLETE"
    )


if __name__ == "__main__":
    main()